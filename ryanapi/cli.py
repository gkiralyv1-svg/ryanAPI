
import argparse
from datetime import date

import requests

try:
    from openpyxl import Workbook, load_workbook
except ImportError:  # pragma: no cover - optional dependency
    Workbook = None
    load_workbook = None


def get_airports():
    url = "https://www.ryanair.com/api/views/locate/3/airports/en/active"

    response = requests.get(url, timeout=30)
    response.raise_for_status()

    data = response.json()
    airports = {}

    for airport in data:
        city = (airport.get("cityCode") or "").lower()
        iata = airport.get("iataCode")

        if not city or not iata:
            continue

        airports.setdefault(city, []).append(iata)

    return airports


def get_ryanair_fares(origin, destination, month, currency="EUR"):
    url = (
        f"https://www.ryanair.com/api/farfnd/v4/"
        f"oneWayFares/{origin}/{destination}/cheapestPerDay"
    )

    params = {
        "outboundMonthOfDate": month,
        "currency": currency,
    }

    response = requests.get(url, params=params, timeout=30)

    if response.status_code != 200:
        print(f"API hiba: {response.status_code}")
        print(response.text)

    response.raise_for_status()

    return response.json()


def save_to_excel(fares, origin, destination):
    if Workbook is None or load_workbook is None:
        return

    filename = "fligths.xlsx"
    sheet_name = f"{date.today()}_{origin}-{destination}"

    try:
        wb = load_workbook(filename)
    except FileNotFoundError:
        wb = Workbook()

    ws = (
        wb[sheet_name]
        if sheet_name in wb.sheetnames
        else wb.create_sheet(sheet_name)
    )

    if ws.max_row == 1:
        ws.append([
            "Dátum",
            "Indulás",
            "Érkezés",
            "Ár",
            "Pénznem",
        ])

    for fare in fares:
        if fare.get("unavailable"):
            continue

        departure = fare["departureDate"][11:16]
        arrival = fare["arrivalDate"][11:16]
        price = fare["price"]["value"]
        currency = fare["price"]["currencyCode"]

        ws.append([
            fare["day"],
            departure,
            arrival,
            price,
            currency,
        ])

    wb.save(filename)


def main():
    parser = argparse.ArgumentParser(
        description="Ryanair járatok és árak lekérdezése"
    )

    parser.add_argument("origin")
    parser.add_argument("destination")
    parser.add_argument("month")

    parser.add_argument(
        "--currency",
        default="EUR",
        help="Pénznem (alapértelmezett: EUR)",
    )

    args = parser.parse_args()

    airports = get_airports()

    origin_key = args.origin.lower()
    destination_key = args.destination.lower()

    if origin_key not in airports:
        parser.error(f"Unknown origin airport '{args.origin}'")

    if destination_key not in airports:
        parser.error(f"Unknown destination airport '{args.destination}'")

    origin_airports = airports[origin_key]
    destination_airports = airports[destination_key]

    print(
        f"Indulási repterek: "
        f"{origin_airports} ({len(origin_airports)} db)"
    )

    print(
        f"Érkezési repterek: "
        f"{destination_airports} ({len(destination_airports)} db)"
    )

    for origin in origin_airports:
        for destination in destination_airports:

            try:
                data = get_ryanair_fares(
                    origin=origin,
                    destination=destination,
                    month=args.month + "-01",
                    currency=args.currency,
                )

            except requests.HTTPError as e:
                print(
                    f"{origin} → {destination}: "
                    f"API hiba ({e.response.status_code})"
                )
                continue

            fares = data.get("outbound", {}).get("fares", [])

            for fare in fares:
                if fare.get("unavailable"):
                    continue

                departure = fare["departureDate"][11:16]
                arrival = fare["arrivalDate"][11:16]
                price = fare["price"]["value"]
                currency = fare["price"]["currencyCode"]

                print(
                    f"{origin} ({args.origin}) → "
                    f"{destination} ({args.destination}) | "
                    f"{fare['day']} "
                    f"{departure} → {arrival} "
                    f"{price:.2f} {currency}"
                )

            save_to_excel(
                fares,
                args.origin,
                args.destination,
            )


if __name__ == "__main__":
    main()
